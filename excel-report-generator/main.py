import sqlite3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


class DatabaseManager:
    """
    كلاس مسؤول عن الاتصال بقاعدة البيانات SQLite
    وإنشاء البيانات الأولية واستخراجها
    """
    def __init__(self, db_name="sales_data.db"):
        self.db_name = db_name

    def setup_database(self):
        """إنشاء جدول المبيعات وإضافة بيانات تجريبية"""
        conn = sqlite3.connect(self.db_name)
        cursor = conn.cursor()
        
        # إنشاء جدول المبيعات في حال لم يكن موجوداً
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS sales (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                product_name TEXT NOT NULL,
                category TEXT,
                quantity INTEGER,
                unit_price REAL,
                total_amount REAL
            )
        ''')
        
        # تنظيف البيانات القديمة لمنع التكرار عند كل تشغيل
        cursor.execute("DELETE FROM sales")
        
        # بيانات تجريبية
        sample_data = [
            ("Laptop Dell XPS", "Electronics", 5, 1200.00, 6000.00),
            ("Wireless Mouse", "Accessories", 15, 25.50, 382.50),
            ("Mechanical Keyboard", "Accessories", 8, 85.00, 680.00),
            ("Monitor 27 Inch", "Electronics", 4, 300.00, 1200.00),
            ("USB-C Hub", "Accessories", 20, 15.00, 300.00)
        ]
        
        cursor.executemany('''
            INSERT INTO sales (product_name, category, quantity, unit_price, total_amount)
            VALUES (?, ?, ?, ?, ?)
        ''', sample_data)
        
        conn.commit()
        conn.close()

    def fetch_sales_data(self):
        """جلب جميع صفوف المبيعات من قاعدة البيانات"""
        conn = sqlite3.connect(self.db_name)
        cursor = conn.cursor()
        cursor.execute("SELECT product_name, category, quantity, unit_price, total_amount FROM sales")
        rows = cursor.fetchall()
        conn.close()
        return rows


class ExcelReportGenerator:
    """
    كلاس مسؤول عن أتمتة وتحويل البيانات إلى تقرير إكسل
    منسق ومصمم بشكل احترافي
    """
    def __init__(self, output_file="Sales_Report.xlsx"):
        self.output_file = output_file
        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        self.ws.title = "Sales Summary"

    def build_report(self, data):
        # 1. كتابة عناوين الأعمدة
        headers = ["Product Name", "Category", "Quantity", "Unit Price ($)", "Total Amount ($)"]
        self.ws.append(headers)

        # 2. كتابة صفوف البيانات
        for row in data:
            self.ws.append(row)

        # 3. إعداد ألوان الخط والخلفية للهيدر (Dark Navy Fill + White Bold Font)
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        for col_num in range(1, len(headers) + 1):
            cell = self.ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # 4. إدراج سطر الإجمالي الكلي ومطبّق عليه دالة SUM الرسمية في إكسل
        last_row = len(data) + 1
        total_row = last_row + 1
        
        self.ws.cell(row=total_row, column=1, value="Total Summary").font = Font(bold=True)
        self.ws.cell(row=total_row, column=5, value=f"=SUM(E2:E{last_row})").font = Font(bold=True, color="1F4E78")

        # 5. التعديل التلقائي لعرض الأعمدة بناءً على طول المحتوى
        for col in self.ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            self.ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # 6. حفظ الملف النهائي
        self.wb.save(self.output_file)
        print(f"[✓] Successfully generated report: {self.output_file}")


# --- نقطة التشغيل الرئيسية ---
if __name__ == "__main__":
    # إنشاء وتجهيز قاعدة البيانات
    db = DatabaseManager()
    db.setup_database()
    sales_data = db.fetch_sales_data()

    # توليد ملف إكسل من البيانات
    reporter = ExcelReportGenerator()
    reporter.build_report(sales_data)