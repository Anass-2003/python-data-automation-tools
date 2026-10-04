
import sqlite3
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

# 1. القراءة من قاعدة البيانات (SQL)
connection = sqlite3.connect("bank.db")
cursor = connection.cursor()

# استعلام جلب جميع بيانات المستخدمين
cursor.execute("SELECT * FROM users")
rows = cursor.fetchall()

connection.close()

# 2. إنشاء ملف Excel جديد منسق
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Bank Users Report"

# كتابة العناوين الرئيسية
headers = ["ID", "Customer Name", "Balance ($)"]
ws.append(headers)

# إضافة البيانات المجلوبة من SQL إلى ملف Excel
for row in rows:
    ws.append([row[0], row[1], row[2]])

# 3. التنسيق والألوان لرأس الجدول (Header Styling)
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
alignment_center = Alignment(horizontal="center", vertical="center")

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = alignment_center

# 4. حفظ ملف Excel
file_name = "bank_users_exported.xlsx"
wb.save(file_name)

print(f"-> Success! Exported {len(rows)} records from SQL to '{file_name}'.")